#!/usr/bin/env python3
"""Versiones editables en Canva de las propuestas de bolsa JDCEL Bq.

Arma un .pptx por diseño (y uno con los nueve) en el que cada pieza es un
elemento aparte: el logo y las figuras van como gráficos vectoriales (SVG con
respaldo PNG), los textos como texto editable y las formas simples como formas.
Canva convierte el .pptx en un diseño editable al importarlo.

Uso:  sh fetch-fonts.sh && python3 canva.py
Necesita: pip install python-pptx fonttools brotli uharfbuzz pyclipper svgelements
          y Node con Playwright (para los PNG de respaldo).

Coordenadas: la cara de la bolsa mide 500 × 600 unidades (1 unidad = 0,5 mm,
25 × 30 cm), igual que en build.mjs.
"""
import hashlib
import io
import json
import math
import os
import re
import shutil
import subprocess

import pyclipper
import uharfbuzz as hb
from fontTools.pens.basePen import BasePen
from fontTools.ttLib import TTFont
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from svgelements import Close, Line, Move
from svgelements import Path as SvgPath

HERE = os.path.dirname(os.path.abspath(__file__))
FONT_DIR = os.path.join(HERE, 'fonts')
OUT = os.path.normpath(os.path.join(HERE, '..', 'canva'))
TMP = os.path.join(HERE, 'out', 'canva-media')

INK = '#111111'  # negro de impresión (también el papel negro)
WHITE = '#FFFFFF'
GLOSS = '#2E2E2E'  # barniz UV sobre negro, solo para verlo en pantalla
GRAY = '#666666'
PX_PER_UNIT = 6  # PNG de respaldo ≈ 300 ppp
EMU_PER_UNIT = 18000  # 0,5 mm
PT_PER_UNIT = 72 / 25.4 / 2
SLOGAN = 'bendecidos para bendecir'


# ------------------------------------------------------------------ fuentes
def _font_files():
    css = open(os.path.join(FONT_DIR, 'fonts.css'), encoding='utf-8').read()
    files = {}
    for subset, body in re.findall(r'/\* ([^*]+) \*/\s*@font-face \{(.*?)\}', css, re.S):
        if subset.strip() != 'latin':
            continue
        family = re.search(r"font-family: '([^']+)'", body).group(1)
        style = re.search(r'font-style: (\w+)', body).group(1)
        src = re.search(r'url\(([^)]+)\)', body).group(1)
        files[(family, style)] = os.path.join(FONT_DIR, src)
    return files


FONT_FILES = _font_files()
_FONTS = {}


def hbfont(family, style='normal', axes=None):
    key = (family, style, tuple(sorted((axes or {}).items())))
    if key not in _FONTS:
        tt = TTFont(FONT_FILES[(family, style)])
        tt.flavor = None
        buf = io.BytesIO()
        tt.save(buf)
        face = hb.Face(hb.Blob(buf.getvalue()))
        font = hb.Font(face)
        if axes:
            font.set_variations(axes)
        _FONTS[key] = (face.upem, font)
    return _FONTS[key]


WORD = dict(family='Archivo', axes={'wght': 900, 'wdth': 125})
MONO = dict(family='Playfair Display', axes={'wght': 900})
SCRIPT = dict(family='Playball')
SERIF = dict(family='EB Garamond', axes={'wght': 400})
SERIF_IT = dict(family='EB Garamond', style='italic', axes={'wght': 400})
SEAL = dict(family='Archivo', axes={'wght': 800, 'wdth': 112})
SANS_B = dict(family='Archivo', axes={'wght': 700, 'wdth': 100})
SANS = dict(family='Archivo', axes={'wght': 400, 'wdth': 100})

# Fuentes de los textos editables: nombre que verá Canva + métrica para ubicarlos.
N_SERIF = dict(font='EB Garamond', m=SERIF)
N_SERIF_IT = dict(font='EB Garamond', italic=True, m=SERIF_IT)
N_SANS_B = dict(font='Archivo', bold=True, m=SANS_B)
N_SANS = dict(font='Archivo', m=SANS)
N_SCRIPT = dict(font='Playball', m=SCRIPT)


# ------------------------------------------------------------------ contornos de texto
class OutlinePen(BasePen):
    """Recibe el contorno de un glifo y lo entrega como path SVG y como polígonos."""

    def __init__(self, xf, steps=14):
        super().__init__(None)
        self.xf, self.steps = xf, steps
        self.d, self.polys, self._poly = [], [], None

    def _moveTo(self, pt):
        self._flush()
        x, y = self.xf(pt)
        self.d.append(f'M{x:.2f} {y:.2f}')
        self._poly = [(x, y)]

    def _lineTo(self, pt):
        x, y = self.xf(pt)
        self.d.append(f'L{x:.2f} {y:.2f}')
        self._poly.append((x, y))

    def _curveToOne(self, p1, p2, p3):
        p0 = self.xf(self._getCurrentPoint())
        a, b, c = self.xf(p1), self.xf(p2), self.xf(p3)
        self.d.append(f'C{a[0]:.2f} {a[1]:.2f} {b[0]:.2f} {b[1]:.2f} {c[0]:.2f} {c[1]:.2f}')
        for i in range(1, self.steps + 1):
            t = i / self.steps
            m = 1 - t
            self._poly.append((
                m ** 3 * p0[0] + 3 * m * m * t * a[0] + 3 * m * t * t * b[0] + t ** 3 * c[0],
                m ** 3 * p0[1] + 3 * m * m * t * a[1] + 3 * m * t * t * b[1] + t ** 3 * c[1],
            ))

    def _closePath(self):
        self.d.append('Z')
        self._flush()

    def _endPath(self):
        self._flush()

    def _flush(self):
        if self._poly and len(self._poly) > 2:
            self.polys.append(self._poly)
        self._poly = None


def _shape(run):
    upem, font = hbfont(run['family'], run.get('style', 'normal'), run.get('axes'))
    buf = hb.Buffer()
    buf.add_str(run['text'])
    buf.guess_segment_properties()
    hb.shape(font, buf, {'kern': True, 'liga': True})
    k = run['size'] / upem
    glyphs, x = [], 0.0
    for info, pos in zip(buf.glyph_infos, buf.glyph_positions):
        glyphs.append((info.codepoint, x + pos.x_offset * k, pos.y_offset * k, pos.x_advance * k))
        x += pos.x_advance * k + run.get('ls', 0)
    return font, k, glyphs, x


def advance(runs):
    return sum(_shape(r)[3] for r in runs)


def outline(runs, x, y, anchor='start', width=None):
    """Texto → (path SVG con curvas, polígonos). `y` es la línea base, como en SVG."""
    shaped = [_shape(r) for r in runs]
    total = sum(s[3] for s in shaped)
    sx = width / total if width else 1.0
    x0 = x - {'start': 0, 'middle': total * sx / 2, 'end': total * sx}[anchor]
    d, polys, start = [], [], 0.0
    for font, k, glyphs, adv in shaped:
        for gid, gx, gy, _ in glyphs:
            pen = OutlinePen(lambda p, ox=start + gx, k=k, gy=gy: (x0 + (ox + p[0] * k) * sx, y - (gy + p[1] * k)))
            font.draw_glyph_with_pen(gid, pen)
            d += pen.d
            polys += pen.polys
        start += adv
    return ''.join(d), polys


def outline_arc(runs, cx, cy, r, top=True):
    """Texto sobre un semicírculo, centrado (como textPath con startOffset 50 %)."""
    shaped = [_shape(rn) for rn in runs]
    total = sum(s[3] for s in shaped)
    s0 = math.pi * r / 2 - total / 2
    d, polys, start = [], [], 0.0
    for font, k, glyphs, adv in shaped:
        for gid, gx, gy, ga in glyphs:
            mid = s0 + start + gx + ga / 2
            phi = math.pi + mid / r if top else math.pi - mid / r
            rho = phi + math.pi / 2 if top else phi - math.pi / 2
            px, py = cx + r * math.cos(phi), cy + r * math.sin(phi)
            cr, sr = math.cos(rho), math.sin(rho)

            def xf(p, k=k, ga=ga, gy=gy, px=px, py=py, cr=cr, sr=sr):
                u, v = p[0] * k - ga / 2, -(gy + p[1] * k)
                return (px + u * cr - v * sr, py + u * sr + v * cr)

            pen = OutlinePen(xf)
            font.draw_glyph_with_pen(gid, pen)
            d += pen.d
            polys += pen.polys
        start += adv
    return ''.join(d), polys


# ------------------------------------------------------------------ geometría
SC = 1000.0


def _ip(polys):
    return [[(round(x * SC), round(y * SC)) for x, y in p] for p in polys if len(p) > 2]


def _fp(paths):
    return [[(x / SC, y / SC) for x, y in p] for p in paths]


def _bool(a, b, op):
    pc = pyclipper.Pyclipper()
    pc.AddPaths(_ip(a), pyclipper.PT_SUBJECT, True)
    if b:
        pc.AddPaths(_ip(b), pyclipper.PT_CLIP, True)
    return _fp(pc.Execute(op, pyclipper.PFT_NONZERO, pyclipper.PFT_NONZERO))


def union(polys):
    return _bool(polys, [], pyclipper.CT_UNION)


def minus(a, b):
    return _bool(a, b, pyclipper.CT_DIFFERENCE)


def inter(a, b):
    return _bool(a, b, pyclipper.CT_INTERSECTION)


def grow(polys, delta):
    pco = pyclipper.PyclipperOffset(2.0, 5.0)
    pco.AddPaths(_ip(union(polys)), pyclipper.JT_ROUND, pyclipper.ET_CLOSEDPOLYGON)
    return _fp(pco.Execute(delta * SC))


def stroke(lines, width, closed=True, round_ends=False):
    """Convierte un trazo en figura rellena (Canva no acepta trazos en SVG)."""
    pco = pyclipper.PyclipperOffset(2.0, 5.0)
    et = pyclipper.ET_CLOSEDLINE if closed else (pyclipper.ET_OPENROUND if round_ends else pyclipper.ET_OPENBUTT)
    for ln in lines:
        pco.AddPath([(round(x * SC), round(y * SC)) for x, y in ln], pyclipper.JT_ROUND, et)
    return _fp(pco.Execute(width / 2 * SC))


def poly_d(polys):
    return ''.join('M' + 'L'.join(f'{x:.2f} {y:.2f}' for x, y in p) + 'Z' for p in polys)


def circle(cx, cy, r, n=None):
    n = n or max(32, int(r * 4))
    return [(cx + r * math.cos(2 * math.pi * i / n), cy + r * math.sin(2 * math.pi * i / n)) for i in range(n)]


def ring(cx, cy, r, w):
    return [circle(cx, cy, r + w / 2), circle(cx, cy, r - w / 2)[::-1]]


def rect(x, y, w, h):
    return [(x, y), (x + w, y), (x + w, y + h), (x, y + h)]


def rrect(x, y, w, h, radii, n=10):
    tl, tr, br, bl = radii if isinstance(radii, (list, tuple)) else [radii] * 4
    pts = []

    def arc(cx, cy, r, a0, a1):
        if r <= 0:
            pts.append((cx, cy))
            return
        for i in range(n + 1):
            a = a0 + (a1 - a0) * i / n
            pts.append((cx + r * math.cos(a), cy + r * math.sin(a)))

    arc(x + tl, y + tl, tl, math.pi, 1.5 * math.pi)
    arc(x + w - tr, y + tr, tr, 1.5 * math.pi, 2 * math.pi)
    arc(x + w - br, y + h - br, br, 0, 0.5 * math.pi)
    arc(x + bl, y + h - bl, bl, 0.5 * math.pi, math.pi)
    return pts


def path_polys(d, tx=0.0, ty=0.0, s=1.0, steps=24):
    """Path SVG (con arcos y curvas) → polilíneas, trasladado y escalado."""
    path = SvgPath(d)
    polys, cur = [], []
    T = lambda q: (tx + q.x * s, ty + q.y * s)  # noqa: E731
    for seg in path:
        if isinstance(seg, Move):
            if len(cur) > 1:
                polys.append(cur)
            cur = [T(seg.end)]
        elif isinstance(seg, Close):
            if len(cur) > 1:
                polys.append(cur)
            cur = []
        elif isinstance(seg, Line):
            cur.append(T(seg.end))
        else:
            cur += [T(seg.point(i / steps)) for i in range(1, steps + 1)]
    if len(cur) > 1:
        polys.append(cur)
    return polys


def bbox(polys, pad=1.0):
    xs = [x for p in polys for x, _ in p]
    ys = [y for p in polys for _, y in p]
    return (min(xs) - pad, min(ys) - pad, max(xs) + pad, max(ys) + pad)


def rng(seed):
    state = [seed]

    def f():
        state[0] = (state[0] * 16807) % 2147483647
        return (state[0] - 1) / 2147483646

    return f


# ------------------------------------------------------------------ piezas de marca
def monogram(cx, cy, s):
    """JJ aproximado (Playfair Display 900) con el corte entre las dos J como hueco real."""
    b1 = cy + s * 0.11
    j = dict(MONO, size=s, text='J')
    _, j1 = outline([j], cx - s * 0.15 + s * 0.03, b1, 'middle')
    d2, j2 = outline([j], cx + s * 0.15 + s * 0.03, b1 + s * 0.2, 'middle')
    j1 = minus(j1, grow(j2, s * 0.0225))
    return poly_d(j1) + d2, j1 + j2


def wordmark(x, y, s, anchor='middle'):
    return outline([dict(WORD, size=s, text='JDCEL')], x, y, anchor)


def bq(x, y, s, anchor='middle'):
    return outline([dict(SCRIPT, size=s, text='Bq')], x, y, anchor)


def cross(cx, cy, h):
    w, t = h * 0.64, h * 0.17
    return union([rect(cx - t / 2, cy - h / 2, t, h), rect(cx - w / 2, cy - h / 2 + h * 0.25, w, t)])


def icon_whatsapp(x, y, s):
    k = s / 24
    bubble = path_polys('M12 2.8a9.2 9.2 0 0 0-7.9 13.9L2.9 21.2l4.6-1.2A9.2 9.2 0 1 0 12 2.8z', x, y, k)
    handset = path_polys('M8.9 7.6c.3-.3.7-.3.9 0l1.2 1.9c.2.3.1.6-.1.9l-.6.7c.7 1.4 1.8 2.5 3.2 3.2l.7-.6c.3-.2.6-.3.9-.1'
                         'l1.9 1.2c.3.2.3.6 0 .9l-.9.9c-.6.6-1.6.8-2.4.4-2.4-1.1-4.3-3-5.4-5.4-.4-.8-.2-1.8.4-2.4z', x, y, k)
    return union(stroke(bubble, 1.9 * k) + handset)


def icon_instagram(x, y, s):
    k = s / 24
    t = lambda pts: [(x + u * k, y + v * k) for u, v in pts]  # noqa: E731
    return union(stroke([t(rrect(3, 3, 18, 18, 5.5))], 1.9 * k) + stroke([t(circle(12, 12, 4.3))], 1.9 * k) + [t(circle(17.2, 6.8, 1.1))])


def icon_pin(x, y, s):
    k = s / 24
    drop = path_polys('M12 21.5s-6.8-6.4-6.8-11.4a6.8 6.8 0 0 1 13.6 0c0 5-6.8 11.4-6.8 11.4z', x, y, k)
    dot = [(x + u * k, y + v * k) for u, v in circle(12, 10.1, 2.4)]
    return union(stroke(drop, 1.9 * k) + stroke([dot], 1.9 * k))


def qr(x, y, size, seed=11):
    n, m, r = 25, size / 25, rng(seed)
    g = [[0] * n for _ in range(n)]

    def finder(ox, oy):
        for i in range(-1, 8):
            for j in range(-1, 8):
                xx, yy = ox + i, oy + j
                if not (0 <= xx < n and 0 <= yy < n):
                    continue
                inside = 0 <= i <= 6 and 0 <= j <= 6
                ring_ = i in (0, 6) or j in (0, 6)
                core = 2 <= i <= 4 and 2 <= j <= 4
                g[yy][xx] = 1 if inside and (ring_ or core) else 2

    finder(0, 0)
    finder(18, 0)
    finder(0, 18)
    for i in range(5):
        for j in range(5):
            g[16 + j][16 + i] = 1 if i in (0, 4) or j in (0, 4) or (i == 2 and j == 2) else 2
    for i in range(8, 17):
        g[6][i] = 1 if i % 2 == 0 else 2
        g[i][6] = 1 if i % 2 == 0 else 2
    cells = []
    for yy in range(n):
        for xx in range(n):
            v = g[yy][xx]
            if v == 0:
                v = 1 if r() < 0.48 else 0
            if v == 1:
                cells.append(rect(x + xx * m, y + yy * m, m, m))
    return union(cells)


def crown(cx, cy, R, seed=7):
    """Corona de espinas de build.mjs, con los trazos convertidos en figuras."""
    r = rng(seed)

    def rad(t, k):
        return R + 10 * math.sin(7 * t + k * 2.1) + 4 * math.sin(3 * t + k * 1.3)

    shapes = []
    for k in range(3):
        pts = [(cx + rad(t, k) * math.cos(t), cy + rad(t, k) * math.sin(t))
               for t in (i / 360 * 2 * math.pi for i in range(360))]
        shapes += stroke([pts], 9 - k * 1.4)
    for i in range(34):
        t = ((i + r() * 0.7) / 34) * math.pi * 2
        k = i % 3
        bx, by = cx + rad(t, k) * math.cos(t), cy + rad(t, k) * math.sin(t)
        outward = i % 3 != 1
        ang = t + (0 if outward else math.pi) + (r() - 0.5) * 1.5
        ln = 17 + r() * 18 if outward else 9 + r() * 7
        w = 3.6 + r() * 1.8
        px, py = -math.sin(ang) * w, math.cos(ang) * w
        shapes.append([(bx + px, by + py), (bx + ln * math.cos(ang), by + ln * math.sin(ang)), (bx - px, by - py)])
    return union(shapes)


def rays(cx, cy, r0):
    tri = []
    for i in range(64):
        a = i / 64 * math.pi * 2 - math.pi / 2
        r1, hw = (170, 3.8) if i % 2 == 0 else (128, 2.7)
        ca, sa = math.cos(a), math.sin(a)
        bx, by = cx + r0 * ca, cy + r0 * sa
        tri.append([(bx - sa * hw, by + ca * hw), (cx + r1 * ca, cy + r1 * sa), (bx + sa * hw, by - ca * hw)])
    return union(tri + ring(cx, cy, r0 - 10, 1.4))


# ------------------------------------------------------------------ elementos
def G(name, layers, export=None):
    """Gráfico vectorial. layers: [(color, d, polígonos)] o [(color, polígonos)]."""
    norm = [(c, *rest) if len(rest) == 2 else (c, poly_d(rest[0]), rest[0]) for c, *rest in layers]
    return dict(kind='graphic', name=name, layers=norm, export=export)


def mono(name, cx, cy, s, color, export=None):
    d, p = monogram(cx, cy, s)
    return G(name, [(color, d, p)], export)


def word(name, x, y, s, color, anchor='middle', export=None):
    d, p = wordmark(x, y, s, anchor)
    return G(name, [(color, d, p)], export)


def script(name, x, y, s, color, anchor='middle', export=None):
    d, p = bq(x, y, s, anchor)
    return G(name, [(color, d, p)], export)


def run(text, spec, size, color, ls=0):
    return dict(text=text, size=size, color=color, ls=ls, **spec)


def T(name, runs, x, y, anchor='middle'):
    return dict(kind='text', name=name, runs=runs, x=x, y=y, anchor=anchor)


def R(name, x, y, w, h, r=0, fill=None, line=None, lw=0):
    return dict(kind='rect', name=name, x=x, y=y, w=w, h=h, r=r, fill=fill, line=line, lw=lw)


def O(name, cx, cy, r, fill=None, line=None, lw=0):
    return dict(kind='oval', name=name, x=cx - r, y=cy - r, w=2 * r, h=2 * r, fill=fill, line=line, lw=lw)


def L(name, x1, y1, x2, y2, color, lw):
    return dict(kind='line', name=name, x1=x1, y1=y1, x2=x2, y2=y2, line=color, lw=lw)


def GROUP(name, children):
    return dict(kind='group', name=name, children=children)


# ------------------------------------------------------------------ los nueve diseños
def designs():
    out = []

    out.append(dict(id='01-corona', title='01 · Corona', bg=WHITE, notes=(
        'Corona de espinas en relieve seco (golpe seco) rodeando el monograma. La corona de esta página es una '
        'imagen de muestra del relieve; el vector para el troquel está en elementos/corona-relieve.svg. '
        'Logo y eslogan en negro mate.'), elements=[
        dict(kind='emboss', name='Corona de espinas (muestra del relieve seco)', polys=crown(250, 250, 122)),
        mono('Logo · monograma JJ (cámbialo por el original)', 250, 250, 146, INK),
        word('Logo · JDCEL', 250, 474, 54, INK),
        script('Logo · Bq', 250, 511, 40, INK),
        T('Eslogan', [run(SLOGAN, N_SERIF, 22, INK)], 250, 548),
    ]))

    status = [rrect(306 + i * 4.2, 123 - h, 2.8, h, 0.6) for i, h in enumerate([4, 6, 8, 10])]
    status += stroke([rrect(330.6, 113.6, 16, 8.8, 2.4)], 1.1) + [rrect(332.5, 115.5, 12, 5, 1), rrect(348.6, 116.4, 1.8, 3.2, 0.6)]
    lock = stroke(path_polys('M245.5 144v-4.2a4.5 4.5 0 0 1 9 0V144'), 1.8, closed=False) + [rrect(242, 143, 16, 12, 2.6)]
    out.append(dict(id='02-pantalla', title='02 · Pantalla', bg=INK, notes=(
        'La bolsa es un celular bloqueado: el monograma hace de reloj, el eslogan de fecha y WhatsApp/Instagram son '
        'los botones. Todo en blanco sobre cartulina negra (serigrafía o hot stamping blanco). Cordón blanco.'), elements=[
        GROUP('Celular', [
            R('Marco', 132, 92, 236, 482, r=42, line=WHITE, lw=4),
            G('Botones laterales', [(WHITE, [rrect(125, 172, 4, 24, 2), rrect(125, 210, 4, 44, 2),
                                             rrect(125, 264, 4, 44, 2), rrect(371, 222, 4, 66, 2)])]),
            R('Isla', 216, 106, 68, 21, r=10.5, fill=WHITE),
            T('Operador', [run('JDCEL', N_SANS_B, 10.5, WHITE)], 153, 122, 'start'),
            G('Señal y batería', [(WHITE, union(status))]),
            G('Candado', [(WHITE, union(lock))]),
            T('Fecha (eslogan)', [run(SLOGAN, N_SERIF_IT, 18, WHITE)], 250, 184),
            mono('Logo · monograma JJ (reloj)', 250, 280, 118, WHITE),
            R('Notificación', 146, 364, 208, 62, r=17, fill=WHITE),
            R('Ícono de la app', 157, 378, 34, 34, r=9, fill=INK),
            mono('Monograma del ícono', 174, 395, 25, WHITE),
            T('Notificación · título', [run('JDCEL ', N_SANS_B, 11, INK), run('Bq', N_SCRIPT, 12.5, INK)], 200, 391, 'start'),
            T('Notificación · hora', [run('ahora', N_SANS, 9.5, GRAY)], 343, 391, 'end'),
            T('Notificación · mensaje', [run('¡Gracias por tu compra!', N_SERIF, 12.5, INK)], 200, 409, 'start'),
            O('Botón WhatsApp', 174, 524, 21, line=WHITE, lw=1.6),
            O('Botón Instagram', 326, 524, 21, line=WHITE, lw=1.6),
            G('Ícono WhatsApp', [(WHITE, icon_whatsapp(163, 513, 22))]),
            G('Ícono Instagram', [(WHITE, icon_instagram(315, 513, 22))]),
            R('Barra de inicio', 205, 557, 90, 5, r=2.5, fill=WHITE),
        ]),
    ]))

    out.append(dict(id='03-resplandor', title='03 · Resplandor', bg=WHITE, notes=(
        'Rayos alrededor del monograma: se leen como luz y como señal. Una sola tinta negra.'), elements=[
        G('Resplandor (rayos)', [(INK, rays(250, 248, 86))], 'resplandor'),
        mono('Logo · monograma JJ (cámbialo por el original)', 250, 248, 110, INK),
        word('Logo · JDCEL', 250, 472, 48, INK),
        script('Logo · Bq', 250, 507, 36, INK),
        T('Eslogan', [run(SLOGAN, N_SERIF, 21, INK)], 250, 546),
    ]))

    x0, y0, g, cw, th, bh, rr, s = 158, 80, 20, 82, 119, 221, 30, 3
    phone = [rrect(x0, y0, cw, th, [rr, s, 0, s]), rrect(x0 + cw + g, y0, cw, th, [s, rr, s, 0]),
             rrect(x0, y0 + th + g, cw, bh, [s, 0, s, rr]), rrect(x0 + cw + g, y0 + th + g, cw, bh, [0, s, rr, s]),
             rrect(153, 130, 3.5, 18, 1.5), rrect(153, 160, 3.5, 30, 1.5), rrect(153, 198, 3.5, 30, 1.5),
             rrect(343.5, 172, 3.5, 46, 1.5)]
    out.append(dict(id='04-cruz-celular', title='04 · Cruz en el celular', bg=WHITE, notes=(
        'La cruz aparece en el espacio blanco que parte la silueta del celular. Una tinta negra; fuelles negros.'), elements=[
        G('Celular con la cruz', [(INK, union(phone))], 'celular-cruz'),
        word('Logo · JDCEL', 232, 500, 46, INK),
        script('Logo · Bq', 340, 503, 36, INK, 'start'),
        T('Eslogan', [run(SLOGAN, N_SERIF, 21, INK)], 250, 542),
    ]))

    d, p = monogram(250, 212, 172)
    logo = union(p + wordmark(250, 410, 60)[1] + bq(250, 456, 46)[1])
    sl = union(outline([dict(SERIF, size=23, text=SLOGAN)], 250, 540, 'middle')[1])
    left, right = [rect(0, 0, 250, 600)], [rect(250, 0, 250, 600)]
    out.append(dict(id='05-luz-y-sombra', title='05 · Luz y sombra', bg=WHITE, notes=(
        'Mitad negra, mitad blanca: todo cambia de color al cruzar la línea (Juan 1:5). El logo y el eslogan partidos '
        'van como gráficos de dos colores.'), elements=[
        R('Mitad negra', 0, 0, 250, 600, fill=INK),
        G('Logo partido (blanco | negro)', [(WHITE, inter(logo, left)), (INK, inter(logo, right))]),
        G('Eslogan partido (blanco | negro)', [(WHITE, inter(sl, left)), (INK, inter(sl, right))]),
    ]))

    pat = []
    for j in range(-1, 9):
        for i in range(-1, 7):
            x, y = 42 + i * 84, 42 + j * 84
            pat += monogram(x, y, 40)[1] + cross(x + 42, y + 42, 17)
    out.append(dict(id='06-monograma', title='06 · Monograma', bg=INK, notes=(
        'Patrón de JJ y cruces en barniz UV brillante sobre laminado mate negro (aquí en gris solo para verlo). '
        'Etiqueta blanca impresa con el logo.'), elements=[
        G('Patrón JJ + cruces (barniz UV)', [(GLOSS, inter(union(pat), [rect(0, 0, 500, 600)]))], 'patron-jj-cruces'),
        GROUP('Etiqueta', [
            R('Etiqueta', 116, 212, 268, 176, fill=WHITE),
            R('Filete', 123.5, 219.5, 253, 161, line=INK, lw=1.2),
            word('Logo · JDCEL', 250, 290, 50, INK),
            script('Logo · Bq', 250, 328, 36, INK),
            L('Línea', 215, 342, 285, 342, INK, 1),
            T('Eslogan', [run(SLOGAN, N_SERIF, 16.5, INK)], 250, 366),
        ]),
    ]))

    out.append(dict(id='07-tipografico', title='07 · Tipográfico', bg=WHITE, notes=(
        'El eslogan como protagonista; Génesis 12:2 es el versículo del que sale «bendecidos para bendecir». '
        'En el fuelle va el QR.'), elements=[
        mono('Logo · monograma JJ (cámbialo por el original)', 70, 132, 66, INK),
        GROUP('Eslogan grande', [
            G('BENDECIDOS', [(INK, *outline([dict(WORD, size=52, text='BENDECIDOS')], 44, 282, width=412))]),
            G('para', [(INK, *outline([dict(SCRIPT, size=74, text='para')], 250, 337, 'middle'))]),
            G('BENDECIR', [(INK, *outline([dict(WORD, size=64, text='BENDECIR')], 44, 414, width=412))]),
        ]),
        L('Línea', 44, 474, 456, 474, INK, 1),
        word('Logo · JDCEL', 44, 518, 26, INK, 'start'),
        script('Logo · Bq', 178, 520, 26, INK, 'start'),
        T('Cita', [run('Génesis 12:2', N_SERIF_IT, 16, INK)], 456, 517, 'end'),
    ]))

    seal = ring(250, 292, 166, 2.6) + ring(250, 292, 157, 1) + ring(250, 292, 111, 1)
    seal += outline_arc([dict(SEAL, size=16.5, ls=3.2, text='BENDECIDOS PARA BENDECIR')], 250, 292, 127, top=True)[1]
    seal += outline_arc([dict(SEAL, size=16.5, ls=3.2, text='JDCEL '), dict(SCRIPT, size=21, text='Bq')], 250, 292, 140, top=False)[1]
    seal += cross(116, 292, 15) + cross(384, 292, 15)
    out.append(dict(id='08-sello', title='08 · Sello', bg=INK, notes=(
        'Sello circular en blanco sobre negro. El mismo sello sirve de sticker para cerrar bolsas y cajas. '
        'Para cambiar las palabras del anillo en Canva, usa un texto con el efecto «Curvar».'), elements=[
        G('Sello (anillos y texto curvo)', [(WHITE, union(seal))], 'sello'),
        mono('Logo · monograma JJ (cámbialo por el original)', 250, 292, 132, WHITE),
    ]))

    out.append(dict(id='09-reverso', title='09 · Reverso', bg=WHITE, notes=(
        'Cara trasera para cualquier opción. Cambia los textos entre corchetes por los datos reales y reemplaza el QR '
        'de ejemplo por uno real (por ejemplo, un enlace wa.me al número de la tienda).'), elements=[
        mono('Logo · monograma JJ (cámbialo por el original)', 250, 152, 74, INK),
        L('Línea', 128, 218, 372, 218, INK, 1),
        G('Ícono WhatsApp', [(INK, icon_whatsapp(128, 243, 26))], 'icono-whatsapp'),
        T('WhatsApp', [run('[tu número]', N_SANS_B, 21, INK)], 168, 262, 'start'),
        G('Ícono Instagram', [(INK, icon_instagram(128, 283, 26))], 'icono-instagram'),
        T('Instagram', [run('[@tu_usuario]', N_SANS_B, 21, INK)], 168, 302, 'start'),
        G('Ícono ubicación', [(INK, icon_pin(128, 323, 26))], 'icono-ubicacion'),
        T('Dirección', [run('[dirección del local]', N_SANS_B, 21, INK)], 168, 342, 'start'),
        G('Código QR (de ejemplo)', [(INK, qr(196, 374, 108, 5))], 'qr-de-ejemplo'),
        T('Texto del QR', [run('escanéalo y escríbenos por WhatsApp', N_SERIF, 15, INK)], 250, 506),
        T('Eslogan', [run(SLOGAN, N_SERIF_IT, 21, INK)], 250, 566),
    ]))
    # ---- ronda 2: a partir de la referencia del cliente (símbolo gigante recortado y logo abajo)
    def page(polys, bleed=6):
        return inter(polys, [rect(-bleed, -bleed, 500 + 2 * bleed, 600 + 2 * bleed)])

    def lockup(color):
        return [word('Logo · JDCEL', 226, 508, 40, color), script('Logo · Bq', 331, 511, 30, color, 'start'),
                T('Eslogan', [run('BENDECIDOS PARA BENDECIR', N_SANS_B, 10.5, color, ls=3.4)], 252, 538)]

    out.append(dict(id='11-monograma-gigante', title='11 · Monograma gigante', round=2, bg=WHITE, notes=(
        'El monograma de la marca en gigante y recortado por los bordes, como el símbolo de la referencia. '
        'Una tinta negra; cordón blanco.'), elements=[
        G('Monograma gigante (recortado)', [(INK, page(monogram(250, 170, 510)[1]))], 'monograma-gigante'),
        *lockup(INK),
    ]))

    rib = lambda r: path_polys(f'M{210 + r} -40V196A{r} {r} 0 0 1 {210 - r} 196', steps=64)  # noqa: E731
    out.append(dict(id='12-jj-en-cinta', title='12 · JJ en cinta', round=2, bg=WHITE, notes=(
        'La J del logo dibujada como una cinta gruesa, dos veces y una dentro de la otra, con curvas como las de la '
        'referencia. Una tinta negra; cordón blanco.'), elements=[
        G('JJ en cinta', [(INK, page(union(stroke(rib(100), 64, closed=False) + stroke(rib(218), 64, closed=False))))], 'jj-en-cinta'),
        *lockup(INK),
    ]))

    cross_lines = [[(250, -260), (250, 330)], [(-200, 118), (700, 118)]]
    fat = lambda w: stroke(cross_lines, w, closed=False, round_ends=True)  # noqa: E731
    out.append(dict(id='13-cruz-en-franjas', title='13 · Cruz en franjas', round=2, bg=WHITE, notes=(
        'Una cruz latina hecha de franjas: las líneas corren de lado a lado y bajan formando curvas en U. '
        'Una tinta negra; cordón blanco.'), elements=[
        G('Cruz en franjas', [(INK, page(union(minus(fat(240), fat(168)) + minus(fat(96), fat(32)))))], 'cruz-en-franjas'),
        *lockup(INK),
    ]))

    a = math.radians(50)
    arcs = [[(250 + r * math.sin(t), 430 - r * math.cos(t)) for t in (-a + 2 * a * i / 80 for i in range(81))]
            for r in (82 + k * 62 for k in range(7))]
    out.append(dict(id='14-senal', title='14 · Señal', round=2, bg=INK, notes=(
        'La señal de wifi en gigante, en blanco sobre la bolsa negra: dice «celulares» y también una bendición que '
        'se comparte. Blanco sobre cartulina negra; cordón negro.'), elements=[
        G('Señal (wifi)', [(WHITE, page(union(stroke(arcs, 34, closed=False, round_ends=True) + [circle(250, 424, 24)])))], 'senal-wifi'),
        *lockup(WHITE),
    ]))

    out.append(dict(id='15-bq-gigante', title='15 · Bq gigante', round=2, bg=INK, notes=(
        'La «Bq» del logo en caligrafía gigante, recortada por los bordes. Blanco sobre cartulina negra; cordón negro.'), elements=[
        G('Bq gigante (recortada)', [(WHITE, page(outline([dict(SCRIPT, size=560, text='Bq')], -40, 345)[1]))], 'bq-gigante'),
        *lockup(WHITE),
    ]))
    return out


# ------------------------------------------------------------------ SVG y PNG
EMBOSS = ('<filter id="e" x="-8%" y="-8%" width="116%" height="116%">'
          '<feGaussianBlur in="SourceAlpha" stdDeviation="1.4" result="b"/>'
          '<feOffset in="b" dx="2" dy="2.4" result="bo"/><feFlood flood-color="#000" flood-opacity=".34"/>'
          '<feComposite in2="bo" operator="in" result="sh"/><feOffset in="b" dx="-1.6" dy="-1.8" result="ho"/>'
          '<feFlood flood-color="#fff" flood-opacity="1"/><feComposite in2="ho" operator="in" result="hl"/>'
          '<feMerge><feMergeNode in="sh"/><feMergeNode in="hl"/><feMergeNode in="SourceGraphic"/></feMerge></filter>')


def svg_doc(layers, box, extra_defs='', wrap=None):
    """SVG apto para Canva: solo rellenos planos, sin trazos, textos, recortes ni filtros."""
    x0, y0, x1, y1 = box
    w, h = x1 - x0, y1 - y0
    body = ''.join(f'<path d="{d}" fill="{c}"/>' for c, d in layers)
    if wrap:
        body = wrap.format(body)
    defs = f'<defs>{extra_defs}</defs>' if extra_defs else ''
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{x0:.2f} {y0:.2f} {w:.2f} {h:.2f}" '
            f'width="{w / 2:.2f}mm" height="{h / 2:.2f}mm">{defs}{body}</svg>')


def prepare_assets(all_designs, media):
    """Escribe el SVG de cada gráfico y renderiza sus PNG de respaldo."""
    manifest, seen = [], {}

    def visit(el):
        if el['kind'] == 'group':
            for c in el['children']:
                visit(c)
            return
        if el['kind'] == 'graphic':
            box = bbox([p for _, _, polys in el['layers'] for p in polys])
            el['box'] = box
            el['svg'] = svg_doc([(c, d) for c, d, _ in el['layers']], box)
        elif el['kind'] == 'emboss':
            box = bbox(el['polys'], pad=6)
            el['box'] = box
            el['svg'] = None
            el['vector'] = svg_doc([(INK, poly_d(el['polys']))], bbox(el['polys']))
            el['render'] = svg_doc([('#FDFDFC', poly_d(el['polys']))], box, EMBOSS, '<g filter="url(#e)">{}</g>')
        else:
            return
        src = el['svg'] or el['render']
        key = hashlib.sha1(src.encode()).hexdigest()[:12]
        el['png'] = os.path.join(media, f'{key}.png')
        if key not in seen:
            seen[key] = True
            x0, y0, x1, y1 = el['box']
            manifest.append(dict(svg=src, png=el['png'], w=math.ceil((x1 - x0) * PX_PER_UNIT), h=math.ceil((y1 - y0) * PX_PER_UNIT)))

    for dsg in all_designs:
        for el in dsg['elements']:
            visit(el)
    path = os.path.join(media, 'manifest.json')
    with open(path, 'w') as fh:
        json.dump(manifest, fh)
    subprocess.run(['node', os.path.join(HERE, 'render_png.mjs'), path], check=True)


def export_elements(all_designs, folder):
    os.makedirs(folder, exist_ok=True)
    done = set()

    def visit(el):
        if el['kind'] == 'group':
            for c in el['children']:
                visit(c)
        elif el.get('export') and el['export'] not in done:
            done.add(el['export'])
            with open(os.path.join(folder, el['export'] + '.svg'), 'w') as fh:
                fh.write(el['svg'])
        elif el['kind'] == 'emboss' and 'corona-relieve' not in done:
            done.add('corona-relieve')
            with open(os.path.join(folder, 'corona-relieve.svg'), 'w') as fh:
                fh.write(el['vector'])

    for dsg in all_designs:
        for el in dsg['elements']:
            visit(el)

    # versiones sueltas del logo en ambos colores, para cualquier fondo
    for name, (d, p) in [('logo-monograma-jj', monogram(0, 0, 146)), ('logo-jdcel', wordmark(0, 0, 54)), ('logo-bq', bq(0, 0, 40))]:
        for suffix, color in [('negro', INK), ('blanco', WHITE)]:
            with open(os.path.join(folder, f'{name}-{suffix}.svg'), 'w') as fh:
                fh.write(svg_doc([(color, d)], bbox(p)))


# ------------------------------------------------------------------ PPTX
def emu(v):
    return Emu(int(round(v * EMU_PER_UNIT)))


def rgb(hex_):
    return RGBColor.from_string(hex_.lstrip('#').upper())


def _plain(shape):
    style = shape._element.find(qn('p:style'))
    if style is not None:
        shape._element.remove(style)


def _paint(shape, el):
    _plain(shape)
    if el.get('fill'):
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(el['fill'])
    else:
        shape.fill.background()
    if el.get('line'):
        shape.line.color.rgb = rgb(el['line'])
        shape.line.width = emu(el['lw'])
    else:
        shape.line.fill.background()


def _svg_blip(slide, pic, svg):
    part = slide.part
    svg_part = Part(part.package.next_image_partname('svg'), 'image/svg+xml', part.package, svg.encode('utf-8'))
    rid = part.relate_to(svg_part, RT.IMAGE)
    blip = pic._element.find('.//' + qn('a:blip'))
    ext_lst = blip.find(qn('a:extLst'))
    if ext_lst is None:
        ext_lst = etree.SubElement(blip, qn('a:extLst'))
    ext = etree.SubElement(ext_lst, qn('a:ext'))
    ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    ns = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
    blip_svg = etree.SubElement(ext, f'{{{ns}}}svgBlip', nsmap={'asvg': ns})
    blip_svg.set(qn('r:embed'), rid)


def add_element(slide, shapes, el):
    kind = el['kind']
    if kind == 'group':
        grp = shapes.add_group_shape()
        grp.name = el['name']
        for child in el['children']:
            add_element(slide, grp.shapes, child)
    elif kind in ('graphic', 'emboss'):
        x0, y0, x1, y1 = el['box']
        pic = shapes.add_picture(el['png'], emu(x0), emu(y0), emu(x1 - x0), emu(y1 - y0))
        pic.name = el['name']
        pic._element.nvPicPr.cNvPr.set('descr', el['name'])
        if el.get('svg'):
            _svg_blip(slide, pic, el['svg'])
    elif kind in ('rect', 'oval'):
        auto = MSO_SHAPE.OVAL if kind == 'oval' else (MSO_SHAPE.ROUNDED_RECTANGLE if el.get('r') else MSO_SHAPE.RECTANGLE)
        shp = shapes.add_shape(auto, emu(el['x']), emu(el['y']), emu(el['w']), emu(el['h']))
        if kind == 'rect' and el.get('r'):
            shp.adjustments[0] = min(0.5, el['r'] / min(el['w'], el['h']))
        shp.name = el['name']
        _paint(shp, el)
    elif kind == 'line':
        ln = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(el['x1']), emu(el['y1']), emu(el['x2']), emu(el['y2']))
        ln.name = el['name']
        _plain(ln)
        ln.line.color.rgb = rgb(el['line'])
        ln.line.width = emu(el['lw'])
    elif kind == 'text':
        runs = el['runs']
        first = runs[0]
        upem, font = hbfont(first['m']['family'], first['m'].get('style', 'normal'), first['m'].get('axes'))
        ext = font.get_font_extents('ltr')
        asc, desc = ext.ascender / upem, -ext.descender / upem
        size = max(r['size'] for r in runs)
        width = sum(advance([dict(r['m'], size=r['size'], text=r['text'], ls=r.get('ls', 0))]) for r in runs) + 40
        left = el['x'] - {'start': 0, 'middle': width / 2, 'end': width}[el['anchor']]
        tb = shapes.add_textbox(emu(left), emu(el['y'] - asc * first['size']), emu(width), emu(size * (asc + desc) * 1.1))
        tb.name = el['name']
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]
        p.alignment = {'start': PP_ALIGN.LEFT, 'middle': PP_ALIGN.CENTER, 'end': PP_ALIGN.RIGHT}[el['anchor']]
        p.line_spacing = 1.0
        for r in runs:
            rn = p.add_run()
            rn.text = r['text']
            rn.font.name = r['font']
            rn.font.size = Pt(round(r['size'] * PT_PER_UNIT, 1))
            rn.font.bold = r.get('bold', False)
            rn.font.italic = r.get('italic', False)
            rn.font.color.rgb = rgb(r['color'])
            if r.get('ls'):
                rn.font._rPr.set('spc', str(round(r['ls'] * PT_PER_UNIT * 100)))


def build_pptx(selected, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = emu(500), emu(600)
    for dsg in selected:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(dsg['bg'])
        for el in dsg['elements']:
            add_element(slide, slide.shapes, el)
        slide.notes_slide.notes_text_frame.text = f"{dsg['title']}. {dsg['notes']}"
    prs.core_properties.title = 'JDCEL Bq · ' + (selected[0]['title'] if len(selected) == 1 else 'propuestas de bolsa')
    prs.save(path)


def _install_static_fonts():
    """LibreOffice no reconoce bien las fuentes variables: instala copias estáticas con nombres simples."""
    from fontTools.varLib import instancer
    dest = os.path.expanduser('~/.fonts/jdcel-canva')
    os.makedirs(dest, exist_ok=True)
    styles = [('Archivo', 'normal', {'wght': 400, 'wdth': 100}, 'Regular'), ('Archivo', 'normal', {'wght': 700, 'wdth': 100}, 'Bold'),
              ('EB Garamond', 'normal', {'wght': 400}, 'Regular'), ('EB Garamond', 'italic', {'wght': 400}, 'Italic'),
              ('Playball', 'normal', None, 'Regular')]
    for family, style, axes, sub in styles:
        tt = TTFont(FONT_FILES[(family, style)])
        tt.flavor = None
        if axes:
            tt = instancer.instantiateVariableFont(tt, axes)
        name = tt['name']
        for rec in list(name.names):
            if rec.nameID in (16, 17, 21, 22, 25):
                name.removeNames(nameID=rec.nameID)
        name.setName(family, 1, 3, 1, 0x409)
        name.setName(sub, 2, 3, 1, 0x409)
        name.setName(f'{family} {sub}', 4, 3, 1, 0x409)
        name.setName(f"{family.replace(' ', '')}-{sub}", 6, 3, 1, 0x409)
        bold, italic = sub == 'Bold', sub == 'Italic'
        tt['OS/2'].usWeightClass = 700 if bold else 400
        tt['OS/2'].fsSelection = (0x20 if bold else 0) | (0x01 if italic else 0) | (0x40 if not (bold or italic) else 0)
        tt['head'].macStyle = (1 if bold else 0) | (2 if italic else 0)
        tt.save(os.path.join(dest, f"{family.replace(' ', '')}-{sub}.ttf"))
    subprocess.run(['fc-cache', '-f', dest], check=False, capture_output=True)


def pdf_copies(folder):
    """Plan B: el mismo diseño en PDF (Canva también lo vuelve editable). Requiere LibreOffice Impress."""
    if not shutil.which('soffice'):
        print('sin LibreOffice: no se generan los PDF')
        return
    _install_static_fonts()
    pptx = sorted(f for f in os.listdir(folder) if re.match(r'\d\d-.*\.pptx$', f))
    subprocess.run(['soffice', '--headless', '--norestore', '--convert-to', 'pdf', '--outdir', os.path.join(folder, 'pdf')]
                   + [os.path.join(folder, f) for f in pptx], check=True, capture_output=True)


def main():
    os.makedirs(OUT, exist_ok=True)
    for name in os.listdir(OUT):  # conserva LEEME.txt
        if name.endswith('.pptx'):
            os.remove(os.path.join(OUT, name))
        elif name in ('pdf', 'elementos'):
            shutil.rmtree(os.path.join(OUT, name))
    shutil.rmtree(TMP, ignore_errors=True)
    os.makedirs(TMP)
    all_designs = designs()
    prepare_assets(all_designs, TMP)
    export_elements(all_designs, os.path.join(OUT, 'elementos'))
    for dsg in all_designs:
        build_pptx([dsg], os.path.join(OUT, f"{dsg['id']}.pptx"))
    build_pptx([d for d in all_designs if d.get('round') != 2], os.path.join(OUT, 'JDCEL-bolsas-9-disenos.pptx'))
    build_pptx([d for d in all_designs if d.get('round') == 2], os.path.join(OUT, 'JDCEL-bolsas-ronda-2.pptx'))
    pdf_copies(OUT)
    shutil.rmtree(TMP)
    print('listo:', OUT)


if __name__ == '__main__':
    main()
